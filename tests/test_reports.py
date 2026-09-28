import os
from datetime import date, timedelta

from openpyxl import load_workbook

from app import reports, seed, services as svc
from app.core import db
from app.models import CommissionLine


def test_relatorio_semanal_gera_arquivos_com_codigos_e_sem_dados_pessoais(app, tmp_path):
    seed.seed_demo()
    end = reports.last_sunday(date.today() + timedelta(days=7))
    ed = reports.generate_weekly(end, str(tmp_path), notes="Semana de estreia do piloto")
    assert os.path.getsize(ed.xlsx_path) > 0 and os.path.getsize(ed.pdf_path) > 0
    wb = load_workbook(ed.xlsx_path)
    assert {"Resumo", "1 Rede de parceiros", "2 Funil de indicações", "3 Financeiro", "9 Critérios do piloto",
            "Retificações", "Dicionário"} <= set(wb.sheetnames)
    text = " ".join(str(c.value) for ws in wb for row in ws.iter_rows() for c in row if c.value is not None)
    assert "P01" in text and "I-001" in text
    for proibido in ("Contábil Horizonte", "Consultoria Mentor", "Metalúrgica", "Distribuidora Aurora"):
        assert proibido not in text
    # reemitir a mesma semana devolve a edição existente (retrato fixo)
    assert reports.generate_weekly(end, str(tmp_path)).id == ed.id


def test_criterios_semaforo_sem_dados_e_com_dados(app):
    crit = {c["chave"]: c for c in reports.evaluate_criteria(date.today())}
    assert crit["calculo"]["status"] == "cinza" and crit["custo"]["status"] == "cinza"
    seed.seed_demo()
    crit = {c["chave"]: c for c in reports.evaluate_criteria(date.today())}
    assert crit["calculo"]["status"] == "amarelo"          # há parcelas ainda não conferidas
    line = CommissionLine.query.filter_by(kind="cliente").first()
    svc.mark_checked(line, True, None, "fin")
    svc.mark_checked(line, False, "Diferença de R$ 0,01", "fin")
    crit = {c["chave"]: c for c in reports.evaluate_criteria(date.today())}
    assert crit["calculo"]["status"] == "vermelho"
    assert crit["juridico"]["status"] in ("vermelho", "amarelo")   # parecer não registrado


def test_retificacao_aparece_na_edicao_seguinte(app, tmp_path):
    seed.seed_demo()
    end1 = reports.last_sunday(date.today() + timedelta(days=7))
    ed1 = reports.generate_weekly(end1, str(tmp_path))
    # correção retroativa: uma indicação da semana 1 é reclassificada para antes do período
    from datetime import datetime
    from app.models import Referral
    ref = Referral.query.filter_by(eligibility="rejeitada").first()
    ref.created_at = datetime.combine(ed1.period_start - timedelta(days=30), datetime.min.time())
    db.session.commit()
    end2 = end1 + timedelta(days=7)
    ed2 = reports.generate_weekly(end2, str(tmp_path))
    wb = load_workbook(ed2.xlsx_path)
    rows = list(wb["Retificações"].iter_rows(values_only=True))
    assert len(rows) >= 2 and rows[1][0] == ed1.edition
    corrigidos = {r[1]: (r[2], r[3]) for r in rows[1:]}
    assert corrigidos["indicacoes_registradas"] == (3, 2)
