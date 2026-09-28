"""Regressões do teste de intrusão: redirecionamento, sessão, duas etapas, links e planilha."""
from datetime import timedelta

import pytest
from openpyxl import load_workbook

from app import reports, seed
from app import security as sec
from app.core import db, now, today
from app.models import Material, Referral, User, UserSession
from app.web.auth import _safe_next


def login(client, email, pw, **kw):
    return client.post("/entrar", data={"email": email, "password": pw}, follow_redirects=False, **kw)


@pytest.fixture()
def world(webapp):
    seed.seed_demo()
    return webapp


def _cookie(client):
    return client.get_cookie("session").value


def _client_with(app, cookie):
    c = app.test_client()
    c.set_cookie("session", cookie)
    return c


# ------------------------------------------------------------------ redirecionamento depois do login
@pytest.mark.parametrize("target", ["//evil.example", "/\\evil.example", "/\t/evil.example", "/\n/evil.example",
                                    "https://evil.example", "javascript:alert(1)", "/\x7f/evil.example", ""])
def test_next_externo_e_recusado(target):
    assert _safe_next(target) is None


@pytest.mark.parametrize("target", ["/painel", "/indicacoes/3", "/gestao/indicacoes?f=abertas", "/%5Cevil.example"])
def test_next_interno_e_aceito(target):
    assert _safe_next(target) == target


def test_login_com_tabulacao_no_next_nao_sai_do_site(world):
    c = world.test_client()
    r = c.post("/entrar?next=/%09/evil.example", data={"email": "parceiro1@demo.test", "password": "senha-segura-123"})
    assert r.status_code == 302
    assert "evil" not in r.headers["Location"]


# ------------------------------------------------------------------ sessão no servidor
def test_sair_invalida_o_cookie_copiado(world):
    c = world.test_client()
    login(c, "parceiro1@demo.test", "senha-segura-123")
    roubado = _cookie(c)
    assert _client_with(world, roubado).get("/painel").status_code == 200
    c.post("/sair")
    assert _client_with(world, roubado).get("/painel").status_code == 302


def test_trocar_a_senha_derruba_as_outras_sessoes(world):
    a, b = world.test_client(), world.test_client()
    login(a, "parceiro1@demo.test", "senha-segura-123")
    login(b, "parceiro1@demo.test", "senha-segura-123")
    copia = _cookie(a)
    r = a.post("/conta/senha", data={"current": "senha-segura-123", "new": "outra-senha-456", "new2": "outra-senha-456"})
    assert r.status_code == 302
    assert a.get("/painel").status_code == 200          # quem trocou continua dentro
    assert b.get("/painel").status_code == 302          # a outra sessão caiu
    assert _client_with(world, copia).get("/painel").status_code == 302   # e o cookie anterior de quem trocou


def test_admin_redefine_senha_ou_desativa_e_a_sessao_cai(world):
    adm, p = world.test_client(), world.test_client()
    login(adm, "gestor@demo.test", "demo-senha-123")
    login(p, "parceiro1@demo.test", "senha-segura-123")
    u = User.query.filter_by(email="parceiro1@demo.test").first()
    adm.post("/gestao/usuarios", data={"action": "reset", "id": u.id, "password": "temporaria-123"})
    assert p.get("/painel").status_code == 302

    login(p, "parceiro1@demo.test", "temporaria-123")
    adm.post("/gestao/usuarios", data={"action": "toggle", "id": u.id})
    assert UserSession.query.filter_by(user_id=u.id).count() == 0


def test_sessao_expira_por_inatividade_e_por_duracao(world):
    c = world.test_client()
    login(c, "parceiro1@demo.test", "senha-segura-123")
    row = UserSession.query.one()
    row.last_seen_at = now() - timedelta(minutes=world.config["SESSION_IDLE_MINUTES"] + 1)
    db.session.commit()
    assert c.get("/painel").status_code == 302
    assert UserSession.query.count() == 0

    login(c, "parceiro1@demo.test", "senha-segura-123")
    row = UserSession.query.one()
    row.created_at = now() - timedelta(hours=world.config["SESSION_MAX_HOURS"] + 1)
    db.session.commit()
    assert c.get("/painel").status_code == 302


def test_rotina_diaria_apaga_sessoes_vencidas(world):
    from app import services as svc
    c = world.test_client()
    login(c, "parceiro1@demo.test", "senha-segura-123")
    UserSession.query.one().last_seen_at = now() - timedelta(days=2)
    db.session.commit()
    assert svc.run_daily_jobs()["sessoes_encerradas"] == 1


def test_cookie_antigo_com_uid_nao_vale_mais(world):
    """Cookies do formato anterior (só o id do usuário) não abrem sessão."""
    c = world.test_client()
    with c.session_transaction() as s:
        s["uid"] = User.query.filter_by(email="gestor@demo.test").first().id
    assert c.get("/gestao/").status_code == 302


# ------------------------------------------------------------------ duas etapas
def _staff_with_2fa():
    u = seed.create_staff("s@t.test", "Staff", "admin", "senha-longa-123")
    u.totp_secret = sec.new_totp_secret()
    db.session.commit()
    return u


def _login_2fa(c, u):
    login(c, "s@t.test", "senha-longa-123")
    assert c.post("/entrar/verificacao", data={"code": sec.totp_now(u.totp_secret)}).status_code == 302


def test_trocar_o_aplicativo_exige_o_codigo_atual(webapp):
    u = _staff_with_2fa()
    antigo = u.totp_secret
    c = webapp.test_client()
    _login_2fa(c, u)
    c.get("/conta/2fa")
    with c.session_transaction() as s:
        novo = s["totp_setup"]
    r = c.post("/conta/2fa", data={"code": sec.totp_now(novo)})             # sem o código atual
    assert b"aplicativo atual incorreto" in r.data
    assert db.session.get(User, u.id).totp_secret == antigo
    r = c.post("/conta/2fa", data={"code": sec.totp_now(novo), "current_code": sec.totp_now(antigo)})
    assert r.status_code == 302
    assert db.session.get(User, u.id).totp_secret == novo


def test_etapa_do_codigo_expira(webapp):
    u = _staff_with_2fa()
    c = webapp.test_client()
    login(c, "s@t.test", "senha-longa-123")
    with c.session_transaction() as s:
        s["pending_at"] = now().timestamp() - 6 * 60
    r = c.post("/entrar/verificacao", data={"code": sec.totp_now(u.totp_secret)})
    assert r.headers["Location"].endswith("/entrar")
    assert c.get("/gestao/").status_code == 302


# ------------------------------------------------------------------ links, cabeçalhos e erros
@pytest.mark.parametrize("url", ["javascript:alert(document.domain)", " JavaScript:alert(1)", "data:text/html,<script>x</script>",
                                 "java\tscript:alert(1)", "//evil.example"])
def test_material_recusa_link_que_nao_e_web(world, url):
    c = world.test_client()
    login(c, "gestor@demo.test", "demo-senha-123")
    n = Material.query.count()
    c.post("/gestao/termos", data={"action": "material", "title": "X", "kind": "link", "url": url})
    assert Material.query.count() == n


def test_material_antigo_com_javascript_nao_vira_link(world):
    db.session.add(Material(title="Velho", kind="link", url="javascript:alert(1)", published=True))
    db.session.commit()
    c = world.test_client()
    login(c, "parceiro1@demo.test", "senha-segura-123")
    html = c.get("/materiais").get_data(as_text=True)
    assert "Velho" in html and 'href="javascript:' not in html


def test_cabecalhos_de_seguranca(webapp):
    r = webapp.test_client().get("/entrar")
    csp = r.headers["Content-Security-Policy"]
    assert "script-src 'self'" in csp and "frame-ancestors 'none'" in csp
    assert "Strict-Transport-Security" not in r.headers        # só com cookie seguro (HTTPS)


def test_hsts_com_https(tmp_path):
    from app import create_app
    app = create_app({"TESTING": True, "SQLALCHEMY_DATABASE_URI": "sqlite:///:memory:", "SECRET_KEY": "x",
                      "SESSION_COOKIE_SECURE": True}, web=True)
    with app.app_context():
        db.create_all()
        assert "max-age=" in app.test_client().get("/entrar").headers["Strict-Transport-Security"]


def test_revisao_com_id_invalido_da_404(world):
    c = world.test_client()
    login(c, "parceiro1@demo.test", "senha-segura-123")
    assert c.get("/revisao?indicacao=abc").status_code == 404


# ------------------------------------------------------------------ planilha
def test_texto_com_igual_nao_vira_formula_na_planilha(app, tmp_path):
    seed.seed_demo()
    ref = Referral.query.filter(Referral.reject_reason.isnot(None)).first()
    ref.reject_reason = '=HYPERLINK("https://evil.example","clique")'
    db.session.commit()
    ed = reports.generate_weekly(reports.last_sunday(today()), str(tmp_path))
    ws = load_workbook(ed.xlsx_path)["2 Funil de indicações"]
    cells = [c for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("=HYPERLINK")]
    assert cells and all(c.data_type == "s" for c in cells)
