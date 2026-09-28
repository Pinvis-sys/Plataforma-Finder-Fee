"""Relatório semanal do piloto (Especificação v2, seção 9).

Fechamento no domingo 23h59, emissão na segunda. Planilha (.xlsx, uma aba por seção, colunas
fixas e dicionário de dados) + resumo em PDF. Os arquivos usam códigos (P01, I-001): nomes e CNPJs
ficam só na plataforma. Correções posteriores aparecem na edição seguinte como retificação.
"""
from __future__ import annotations

import os
from datetime import date, datetime, timedelta
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from . import engine as eng
from .core import db, now
from .models import (CommissionLine, Contract, DuplicateAttempt, Partner, Receipt, Referral, ReportEdition,
                     ReviewRequest, SurveyResponse, setting)
from .rules import active_ruleset, holidays, recruit_cap_for_year

VINHO, CARMIM, LARANJA = "#550A19", "#820B25", "#DF581B"
STATUS_LABEL = {"verde": "Dentro da meta", "amarelo": "Atenção", "vermelho": "Fora da meta", "cinza": "Sem dados"}
STATUS_COLOR = {"verde": "#2E7D32", "amarelo": "#B8860B", "vermelho": "#B3261E", "cinza": "#8A8A8A"}


def last_sunday(d: date) -> date:
    return d - timedelta(days=(d.weekday() + 1) % 7)


def week_bounds(end: date) -> tuple[datetime, datetime]:
    start = datetime.combine(end - timedelta(days=6), datetime.min.time())
    stop = datetime.combine(end + timedelta(days=1), datetime.min.time())
    return start, stop


def edition_id(end: date) -> str:
    iso = end.isocalendar()
    return f"{iso[0]}-W{iso[1]:02d}"


def _pilot_start(rules: dict) -> datetime:
    return datetime.combine(date.fromisoformat(rules["pilot_start"]), datetime.min.time())


def _bd(a, b, hol) -> int:
    return eng.business_days_between(a.date(), b.date(), hol)


# ------------------------------------------------------------------ contadores
def counters(start: datetime, stop: datetime) -> dict:
    n = lambda q: q.count()  # noqa: E731
    recv = (db.session.query(Receipt).filter(Receipt.created_at >= start, Receipt.created_at < stop).all())
    return {
        "candidaturas": n(Partner.query.filter(Partner.application_at >= start, Partner.application_at < stop)),
        "parceiros_ativados": n(Partner.query.filter(Partner.activated_at >= start, Partner.activated_at < stop)),
        "indicacoes_registradas": n(Referral.query.filter(Referral.created_at >= start, Referral.created_at < stop)),
        "indicacoes_aceitas": n(Referral.query.filter(Referral.accepted_at >= start, Referral.accepted_at < stop)),
        "indicacoes_rejeitadas": n(Referral.query.filter(Referral.eligibility == "rejeitada", Referral.reviewed_at >= start,
                                                        Referral.reviewed_at < stop)),
        "contratos": n(Contract.query.filter(Contract.created_at >= start, Contract.created_at < stop,
                                             Contract.cancelled_at.is_(None))),
        "recebido_clientes": str(sum((r.amount for r in recv), Decimal("0"))),
        "comissoes_pagas": n(CommissionLine.query.filter(CommissionLine.state == "paga", CommissionLine.paid_at >= start.date(),
                                                         CommissionLine.paid_at < stop.date())),
    }


# ------------------------------------------------------------------ critérios do piloto
def evaluate_criteria(end: date) -> list[dict]:
    rules = active_ruleset().rules
    hol = holidays(rules)
    sla = int(rules["sla_review_days"])
    pilot_end = date.fromisoformat(rules["pilot_end"])
    ps = _pilot_start(rules)
    out: list[dict] = []

    def add(group, key, name, measure, target, value, status):
        out.append({"grupo": group, "chave": key, "criterio": name, "como_medir": measure, "meta": target,
                    "valor": value, "status": status})

    G = "Eliminatório"
    lines = CommissionLine.query.filter(CommissionLine.kind != "ajuste").all()
    if not lines:
        add(G, "calculo", "Cálculo correto", "Conferência de 100% das parcelas contra a planilha",
            "Zero divergência sem explicação e nenhuma duplicidade", "Sem parcelas ainda", "cinza")
    else:
        unchecked = sum(1 for l in lines if not l.checked and not l.divergence_note)
        diverg = sum(1 for l in lines if l.divergence_note)
        st = "vermelho" if diverg else ("verde" if unchecked == 0 else "amarelo")
        add(G, "calculo", "Cálculo correto", "Conferência de 100% das parcelas contra a planilha",
            "Zero divergência sem explicação e nenhuma duplicidade",
            f"{len(lines)} parcelas; {unchecked} a conferir; {diverg} com divergência", st)

    paid_ok = [l for l in lines if l.kind == "cliente" and l.state == "paga" and l.receipt
               and l.cycle == eng.next_month(l.receipt.received_at)]
    any_paid = any(l.state == "paga" for l in lines)
    has_contract = Contract.query.filter(Contract.cancelled_at.is_(None)).count() > 0
    if paid_ok:
        st, val = "verde", "1 ciclo completo (contrato, recebimento, NF e pagamento no mês seguinte)"
    elif any_paid:
        st, val = "amarelo", "Pagamento ocorreu fora do mês subsequente"
    elif has_contract:
        st, val = ("vermelho" if end > pilot_end else "amarelo"), "Contrato assinado, ciclo ainda não concluído"
    else:
        st, val = "cinza", "Sem contratos ainda"
    add(G, "ciclo", "Ciclo financeiro completo", "Contrato percorre assinatura, recebimento, NF e pagamento no mês seguinte",
        "1 ciclo sem intervenção manual fora do previsto", val, st)

    disputes = DuplicateAttempt.query.filter_by(disputed=True, resolved=False).count()
    any_ref = Referral.query.count() > 0
    add(G, "atribuicao", "Atribuição", "Disputas ou duplicidades de indicação", "Nenhuma disputa sem solução pela regra de data e hora",
        f"{disputes} disputa(s) sem solução", "vermelho" if disputes else ("verde" if any_ref else "cinza"))

    opinion = setting("legal_opinion_date")
    pending_consent = Referral.query.filter(Referral.eligibility == "aceita", Referral.stage != "aguardando_contato",
                                            Referral.consent_status != "registrado").count()
    if pending_consent or (has_contract and not opinion):
        st = "vermelho"
    elif opinion:
        st = "verde"
    else:
        st = "amarelo"
    add(G, "juridico", "Conformidade jurídica", "Parecer antes do primeiro contrato; consentimento LGPD antes do contato",
        "100%", f"Parecer: {opinion or 'não registrado'}; consentimentos pendentes com contato: {pending_consent}", st)

    alerts = _frequency_count(rules, end)
    add(G, "olheiro", "Hipótese do olheiro eventual", "Sinal observado em cada indicação; frequência por parceiro",
        "Nenhum padrão de prospecção contínua; indicações ligadas a um sinal",
        f"{alerts} parceiro(s) com alta recorrência", "vermelho" if alerts else ("verde" if any_ref else "cinza"))

    net = sum((l.base_net or Decimal("0") for l in lines if l.kind == "cliente"), Decimal("0"))
    cost = sum((l.amount for l in lines if l.kind in ("cliente", "bonus")), Decimal("0"))
    ratio = (cost / net) if net else None
    limit = rules.get("channel_cost_limit")
    if ratio is None:
        st, val = "cinza", "Sem receita líquida ainda"
    elif limit is None:
        st, val = "cinza", f"{ratio:.1%} da receita líquida (limite a definir com a precificação)"
    else:
        st, val = ("verde" if ratio <= Decimal(str(limit)) else "vermelho"), f"{ratio:.1%} da receita líquida"
    add(G, "custo", "Custo do canal", "Comissão mais bônus sobre a receita líquida", "Dentro do que a margem suporta", val, st)

    R = "Resultado"
    active = Partner.query.filter_by(status="ativo").all()
    with_ref = sum(1 for p in active if Referral.query.filter_by(partner_id=p.id).count() > 0)
    if not active:
        add(R, "ativacao", "Ativação", "Parceiros aprovados com ao menos 1 indicação", "60% ou mais", "Sem parceiros ativos", "cinza")
    else:
        r_ = with_ref / len(active)
        add(R, "ativacao", "Ativação", "Parceiros aprovados com ao menos 1 indicação", "60% ou mais",
            f"{with_ref} de {len(active)} ({r_:.0%})", "verde" if r_ >= .6 else ("amarelo" if r_ >= .4 else "vermelho"))

    reg = Referral.query.count()
    acc = Referral.query.filter(Referral.accepted_at.isnot(None)).count()
    add(R, "volume", "Volume", "Indicações registradas e aceitas", "4 ou mais registradas e 2 ou mais aceitas",
        f"{reg} registradas; {acc} aceitas", "verde" if reg >= 4 and acc >= 2 else ("amarelo" if reg or acc else "vermelho"))

    judged = Referral.query.filter(Referral.accepted_at.isnot(None), Referral.fits_target.isnot(None)).all()
    if judged:
        q_ = sum(1 for r in judged if r.fits_target) / len(judged)
        add(R, "qualidade", "Qualidade", "Aderência ao cliente ideal e à faixa-alvo", "75% ou mais das aceitas",
            f"{q_:.0%} das aceitas", "verde" if q_ >= .75 else ("amarelo" if q_ >= .5 else "vermelho"))
    else:
        add(R, "qualidade", "Qualidade", "Aderência ao cliente ideal e à faixa-alvo", "75% ou mais das aceitas", "Sem aceitas avaliadas", "cinza")

    n_contracts = Contract.query.filter(Contract.cancelled_at.is_(None)).count()
    add(R, "conversao", "Conversão", "Contratos assinados sobre indicações aceitas", "Ao menos 1 contrato",
        f"{n_contracts} contrato(s) em {acc} aceitas", "verde" if n_contracts else ("vermelho" if end > pilot_end else "amarelo"))

    decided = Referral.query.filter(Referral.reviewed_at.isnot(None)).all()
    v1 = [(_bd(r.created_at, r.reviewed_at, hol) <= sla) for r in decided]
    contacted = Referral.query.filter(Referral.first_contact_at.isnot(None), Referral.accepted_at.isnot(None)).all()
    v2 = [(_bd(r.accepted_at, r.first_contact_at, hol) <= int(rules["first_contact_days"])) for r in contacted]
    if v1 or v2:
        p1_ = sum(v1) / len(v1) if v1 else None
        p2_ = sum(v2) / len(v2) if v2 else None
        vals = [x for x in (p1_, p2_) if x is not None]
        st = "verde" if all(x >= .9 for x in vals) else ("amarelo" if all(x >= .7 for x in vals) else "vermelho")
        add(R, "velocidade", "Velocidade", "Validação e primeiro contato",
            f"Validação em até {sla} dias úteis em 90% dos casos; primeiro contato em até {rules['first_contact_days']} dias úteis",
            f"Validação no prazo: {'—' if p1_ is None else f'{p1_:.0%}'}; primeiro contato no prazo: {'—' if p2_ is None else f'{p2_:.0%}'}", st)
    else:
        add(R, "velocidade", "Velocidade", "Validação e primeiro contato", f"Validação em até {sla} dias úteis em 90% dos casos", "Sem dados", "cinza")

    recruited = Partner.query.filter(Partner.invited_by_id.isnot(None), Partner.status != "rejeitado").all()
    if not recruited:
        add(R, "camada2", "Camada 2", "Recrutado conclui candidatura, aprovação, workshop e termos, e entende o aviso do bônus",
            "Fluxo completo e nenhum 'não sabia'", "Sem recrutados", "cinza")
    else:
        done = [p for p in recruited if p.status == "ativo" and p.bonus_notice_accepted_at]
        add(R, "camada2", "Camada 2", "Recrutado conclui candidatura, aprovação, workshop e termos, e entende o aviso do bônus",
            "Fluxo completo e nenhum 'não sabia'", f"{len(done)} de {len(recruited)} concluíram o fluxo", "verde" if len(done) == len(recruited) else "amarelo")

    surveys = SurveyResponse.query.all()
    if surveys:
        avg = sum(s.score for s in surveys) / len(surveys)
        doubts = sum(1 for s in surveys if s.calc_doubt)
        add(R, "experiencia", "Experiência do parceiro", "Mini-pesquisa final e dúvidas recorrentes",
            "Nota média 8 ou mais e nenhuma dúvida recorrente sobre o cálculo", f"Média {avg:.1f}; {doubts} dúvida(s) sobre o cálculo",
            "verde" if avg >= 8 and not doubts else ("amarelo" if avg >= 7 else "vermelho"))
    else:
        add(R, "experiencia", "Experiência do parceiro", "Mini-pesquisa final e dúvidas recorrentes", "Nota média 8 ou mais", "Pesquisa não aplicada", "cinza")

    late = Referral.query.filter_by(delayed_capacity=True).count()
    add(R, "capacidade", "Capacidade", "Atendimentos atrasados por falta de capacidade dos sócios", "Nenhum",
        f"{late} atendimento(s) atrasado(s)", "vermelho" if late else "verde")
    return out


def _frequency_count(rules: dict, end: date) -> int:
    from .services import frequency_alerts
    return len(frequency_alerts(rules, end))


# ------------------------------------------------------------------ coleta por seção
def collect(end: date, notes: str = "") -> dict:
    rules = active_ruleset().rules
    hol = holidays(rules)
    start, stop = week_bounds(end)
    ps = _pilot_start(rules)
    week, cum = counters(start, stop), counters(ps, stop)
    code_p = lambda pid: db.session.get(Partner, pid).code  # noqa: E731
    S: dict[str, dict] = {}

    partners = Partner.query.order_by(Partner.code).all()
    thirty = stop - timedelta(days=30)
    rows = []
    for p in partners:
        last_ref = (Referral.query.filter_by(partner_id=p.id).order_by(Referral.created_at.desc()).first())
        rows.append([p.code, p.status, "Recrutado" if p.invited_by_id else "Inicial", code_p(p.invited_by_id) if p.invited_by_id else "",
                     p.application_at.date().isoformat(), p.approved_at.date().isoformat() if p.approved_at else "",
                     "sim" if p.workshop_at else "não",
                     "parado" if p.status == "ativo" and (not last_ref or last_ref.created_at < thirty) else ""])
    if rules["pilot_mode"]:
        rows.append(["", "", "", "", "", "", "Recrutamentos usados/limite", f"{_recruit_used()}/{rules['pilot_recruit_limit']}"])
    S["1 Rede de parceiros"] = {"headers": ["Parceiro", "Situação", "Origem", "Convidado por", "Candidatura", "Aprovação", "Workshop", "Observação"], "rows": rows}

    refs = Referral.query.order_by(Referral.code).all()
    rows = []
    for r in refs:
        rows.append([r.code, code_p(r.partner_id), r.created_at.date().isoformat(), r.eligibility,
                     "espera" if r.waiting else "", r.stage, r.reject_reason or "", r.lost_reason or "",
                     _bd(r.created_at, r.reviewed_at, hol) if r.reviewed_at else "",
                     _bd(r.accepted_at, r.first_contact_at, hol) if r.accepted_at and r.first_contact_at else ""])
    S["2 Funil de indicações"] = {"headers": ["Indicação", "Parceiro", "Registro", "Elegibilidade", "Espera", "Etapa comercial",
                                              "Motivo da rejeição", "Motivo da perda", "Dias úteis até a decisão", "Dias úteis até o 1º contato"], "rows": rows}
    dups = DuplicateAttempt.query.count()
    S["2 Funil de indicações"]["rows"].append(["", "", "", "", "", "", "", "Duplicidades barradas (acumulado)", dups, ""])

    lines = CommissionLine.query.order_by(CommissionLine.id).all()
    rows = []
    for l in lines:
        rows.append([l.id, l.kind, code_p(l.beneficiary_id), l.referral.code if l.referral else "", l.state,
                     float(l.amount), l.receipt.received_at.isoformat() if l.receipt else "", l.cycle.isoformat() if l.cycle else "",
                     "sim" if l.checked else "não", l.divergence_note or ""])
    S["3 Financeiro"] = {"headers": ["Linha", "Tipo", "Beneficiário", "Indicação", "Estado", "Valor (R$)", "Recebimento do cliente",
                                     "Ciclo de pagamento", "Conferida", "Divergência"], "rows": rows}

    open_reviews = ReviewRequest.query.filter_by(status="aberto").count()
    S["4 Qualidade operacional"] = {"headers": ["Indicador", "Valor"], "rows": [
        ["Pedidos de revisão em aberto", open_reviews], ["Pedidos de revisão (total)", ReviewRequest.query.count()],
        ["Linhas com divergência de cálculo", sum(1 for l in lines if l.divergence_note)],
        ["Ajustes manuais", sum(1 for l in lines if l.kind == "ajuste")],
        ["Indicações aguardando decisão", Referral.query.filter(Referral.eligibility.in_(("recebida", "em_validacao"))).count()]]}

    from .services import frequency_alerts
    S["5 Risco e conformidade"] = {"headers": ["Indicador", "Parceiro/Indicação", "Valor"], "rows":
        [["Alta recorrência (12 meses)", p.code, n] for p, n in frequency_alerts(rules, end)]
        + [["Conflito de interesse sinalizado", p.code, "sim"] for p in partners if p.conflict_notes]
        + [["Consentimento LGPD pendente", r.code, r.stage] for r in refs if r.eligibility == "aceita" and r.consent_status == "pendente"]
        + [["CNPJ sem confirmação de situação ativa", p.code, p.status] for p in partners if p.status == "ativo" and not p.cnpj_verified]}

    in_service = Referral.query.filter(Referral.eligibility == "aceita", Referral.waiting.is_(False),
                                       Referral.stage.notin_(("perdida",))).count()
    S["6 Capacidade"] = {"headers": ["Indicador", "Valor"], "rows": [
        ["Relações em atendimento por indicação (Finder Fee)", in_service],
        ["Indicações em espera (fila interna)", Referral.query.filter_by(waiting=True, eligibility="aceita").count()],
        ["Capacidade estimada total (estimativa, todos os canais)", setting("capacity_estimate", "40 a 50 relações")]]}

    prev_ed = ReportEdition.query.filter(ReportEdition.period_end < end).order_by(ReportEdition.period_end.desc()).first()
    prev_week = (prev_ed.metrics or {}).get("week", {}) if prev_ed else {}
    comp = [[k, week[k], prev_week.get(k, ""), cum[k]] for k in week]
    S["7 Comparativo"] = {"headers": ["Indicador", "Esta semana", "Semana anterior", "Acumulado do piloto"], "rows": comp}

    S["8 Observações do gestor"] = {"headers": ["Observações"], "rows": [[notes or ""]]}
    crit = evaluate_criteria(end)
    S["9 Critérios do piloto"] = {"headers": ["Grupo", "Critério", "Como medir", "Meta", "Situação atual", "Semáforo"],
                                  "rows": [[c["grupo"], c["criterio"], c["como_medir"], c["meta"], c["valor"], STATUS_LABEL[c["status"]]] for c in crit]}

    retif = []
    if prev_ed:
        ps_, pe_ = week_bounds(prev_ed.period_end)
        recomputed = counters(ps_, pe_)
        for k, v in (prev_ed.metrics or {}).get("week", {}).items():
            if str(recomputed.get(k)) != str(v):
                retif.append([prev_ed.edition, k, v, recomputed.get(k)])
    S["Retificações"] = {"headers": ["Edição corrigida", "Indicador", "Valor publicado", "Valor corrigido"], "rows": retif}
    S["Dicionário"] = {"headers": ["Aba", "Coluna/indicador", "Significado"], "rows": DICTIONARY}
    return {"sections": S, "week": week, "cum": cum, "criteria": crit, "notes": notes}


def _recruit_used() -> int:
    return Partner.query.filter(Partner.invited_by_id.isnot(None), Partner.status != "rejeitado").count()


DICTIONARY = [
    ["Todas", "Parceiro (P01…) e Indicação (I-001…)", "Códigos estáveis. Nomes e CNPJs ficam apenas na plataforma."],
    ["1 Rede de parceiros", "Situação", "candidato, em_analise, aprovado (falta workshop/termos), ativo, rejeitado, suspenso"],
    ["1 Rede de parceiros", "Observação = parado", "Parceiro ativo sem indicação há mais de 30 dias"],
    ["2 Funil de indicações", "Elegibilidade", "recebida, em_validacao, aceita, rejeitada, expirada"],
    ["2 Funil de indicações", "Espera", "Aceita, mas aguardando início do atendimento (fila interna)"],
    ["2 Funil de indicações", "Dias úteis", "Contados sem fins de semana e feriados cadastrados"],
    ["3 Financeiro", "Tipo", "cliente (comissão), bonus (Camada 2), ajuste (manual)"],
    ["3 Financeiro", "Estado", "aguardando_condicao, aprovada, programada, paga, ajustada, cancelada"],
    ["3 Financeiro", "Conferida", "Conferência de 100% das parcelas contra a planilha (critério eliminatório)"],
    ["7 Comparativo", "Indicador", "Contagens da semana (segunda a domingo) contra a anterior e o acumulado do piloto"],
    ["9 Critérios do piloto", "Semáforo", "Dentro da meta, Atenção, Fora da meta, Sem dados"],
    ["Retificações", "Valor corrigido", "Diferença entre o que foi publicado na edição anterior e o que os dados mostram hoje"],
]


# ------------------------------------------------------------------ arquivos
def _write_xlsx(path: str, data: dict, edition: str, end: date):
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumo"
    head = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor=VINHO.lstrip("#"))
    ws.append(["Relatório semanal do piloto — Finder Fee All Targets"])
    ws["A1"].font = Font(bold=True, size=14, color=VINHO.lstrip("#"))
    ws.append([f"Edição {edition} · fechamento em {end:%d/%m/%Y} 23:59 · arquivo com códigos, sem nomes ou CNPJs"])
    ws.append([])
    ws.append(["Indicador", "Esta semana", "Acumulado"])
    for c in ws[4]:
        c.font, c.fill = head, fill
    for k in data["week"]:
        ws.append([k.replace("_", " ").capitalize(), data["week"][k], data["cum"][k]])
    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 20
    ws.column_dimensions["C"].width = 20
    for name, sec in data["sections"].items():
        s = wb.create_sheet(name[:31])
        s.append(sec["headers"])
        for c in s[1]:
            c.font, c.fill = head, fill
        for row in sec["rows"]:
            s.append(row)
            for c in s[s.max_row]:
                # texto começando com "=" viraria fórmula no Excel (texto digitado pode vir do parceiro)
                if isinstance(c.value, str) and c.value.startswith("="):
                    c.data_type = "s"
        s.freeze_panes = "A2"
        for i, h in enumerate(sec["headers"], start=1):
            width = max([len(str(h))] + [len(str(r[i - 1])) for r in sec["rows"] if len(r) >= i] + [10])
            s.column_dimensions[s.cell(1, i).column_letter].width = min(width + 2, 60)
    wb.save(path)


def _write_pdf(path: str, data: dict, edition: str, end: date):
    styles = getSampleStyleSheet()
    h1 = ParagraphStyle("h1", parent=styles["Title"], textColor=colors.HexColor(VINHO), fontName="Helvetica-Bold", fontSize=16, alignment=0)
    body = ParagraphStyle("b", parent=styles["BodyText"], fontSize=9, leading=12, textColor=colors.HexColor("#2D2D2D"))
    small = ParagraphStyle("s", parent=body, fontSize=8, textColor=colors.HexColor("#5C5C5C"))
    doc = SimpleDocTemplate(path, pagesize=A4, leftMargin=16 * mm, rightMargin=16 * mm, topMargin=14 * mm, bottomMargin=14 * mm,
                            title=f"Relatório semanal {edition}")
    els = [Paragraph("Relatório semanal do piloto — Finder Fee", h1),
           Paragraph(f"Edição {edition} · fechamento em {end:%d/%m/%Y} às 23h59 · códigos no lugar de nomes e CNPJs", small), Spacer(1, 6)]
    rows = [["Indicador", "Semana", "Acumulado"]] + [[k.replace("_", " ").capitalize(), str(data["week"][k]), str(data["cum"][k])] for k in data["week"]]
    t = Table(rows, colWidths=[90 * mm, 40 * mm, 40 * mm])
    t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(VINHO)), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                           ("FONTSIZE", (0, 0), (-1, -1), 8.5), ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F5F0F0")]),
                           ("GRID", (0, 0), (-1, -1), .25, colors.HexColor("#E8E0E0")), ("ALIGN", (1, 0), (-1, -1), "RIGHT")]))
    els += [t, Spacer(1, 8), Paragraph("<b>Critérios do piloto</b>", body)]
    crows = [["Critério", "Situação atual", "Semáforo"]]
    for c in data["criteria"]:
        crows.append([Paragraph(f"{c['criterio']}<br/><font size=7 color='#5C5C5C'>{c['grupo']} · meta: {c['meta']}</font>", body),
                      Paragraph(c["valor"], body), STATUS_LABEL[c["status"]]])
    ct = Table(crows, colWidths=[72 * mm, 68 * mm, 30 * mm], repeatRows=1)
    style = [("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(VINHO)), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
             ("FONTSIZE", (0, 0), (-1, -1), 8), ("GRID", (0, 0), (-1, -1), .25, colors.HexColor("#E8E0E0")), ("VALIGN", (0, 0), (-1, -1), "MIDDLE")]
    for i, c in enumerate(data["criteria"], start=1):
        style += [("BACKGROUND", (2, i), (2, i), colors.HexColor(STATUS_COLOR[c["status"]])), ("TEXTCOLOR", (2, i), (2, i), colors.white)]
    ct.setStyle(TableStyle(style))
    els += [ct, Spacer(1, 8)]
    if data.get("notes"):
        els += [Paragraph("<b>Observações do gestor</b>", body), Paragraph(data["notes"], body)]
    doc.build(els)


def generate_weekly(end: date, out_dir: str, notes: str = "", force: bool = False) -> ReportEdition:
    """Cada edição é um retrato fixo: reemitir a mesma semana exige force=True. Correções de dados
    aparecem na edição seguinte, na aba Retificações."""
    os.makedirs(out_dir, exist_ok=True)
    ed_id = edition_id(end)
    existing = ReportEdition.query.filter_by(edition=ed_id).first()
    if existing and not force:
        return existing
    data = collect(end, notes)
    xlsx = os.path.join(out_dir, f"relatorio_piloto_{ed_id}.xlsx")
    pdf = os.path.join(out_dir, f"relatorio_piloto_{ed_id}.pdf")
    _write_xlsx(xlsx, data, ed_id, end)
    _write_pdf(pdf, data, ed_id, end)
    ed = ReportEdition.query.filter_by(edition=ed_id).first()
    if ed is None:
        ed = ReportEdition(edition=ed_id, period_start=end - timedelta(days=6), period_end=end)
        db.session.add(ed)
    ed.metrics = {"week": data["week"], "cum": data["cum"], "criteria": [{k: c[k] for k in ("chave", "status", "valor")} for c in data["criteria"]]}
    ed.xlsx_path, ed.pdf_path, ed.generated_at = xlsx, pdf, now()
    db.session.commit()
    return ed
